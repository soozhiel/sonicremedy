#!/usr/bin/env python3
"""Build data/expressions.json from TAB 05 — HUMAN EXPRESSION LIBRARY of the master Sheet.

Source of truth: the master Sheet 'SONIC REMEDY — MASTER SONG INDEX' (same export as
build_songs.py: File > Download > Microsoft Excel, or a Drive export of the same file).

Usage:  python3 scripts/build_expressions.py <master.xlsx> [data/expressions.json] [data/songs.json]

Only rows whose Expression ID looks like EXP-### are read (the notes row is skipped). Every
linked song ID must exist in TAB 02 and in data/songs.json, otherwise the build fails: TAB 02
stays the authority and nothing is linked that the song records do not support. Rows with no
link are kept (songs = []) so the engine knows the phrase was considered, but they never add
a recommendation. Public data only: sources, notes and match-basis text stay in the Sheet.
"""
import sys, json, re, datetime
from openpyxl import load_workbook

ID_RE = re.compile(r'^EXP-\d{3,}$')
SONG_RE = re.compile(r'SR-F\d{2}-T\d{2}')
CONF = {'high', 'medium', 'low-medium', 'low', 'n/a'}

def tab(wb, prefix):
    for n in wb.sheetnames:
        if n.upper().startswith(prefix):
            return wb[n]
    raise SystemExit(f'missing tab {prefix}')

def rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else '' for h in next(it)]
    for r in it:
        if r and any(c not in (None, '') for c in r):
            yield {h: ('' if v is None else str(v).strip()) for h, v in zip(head, r)}

def col(r, start):
    for k in r:
        if k.lower().startswith(start.lower()):
            return r[k]
    raise SystemExit(f'missing column {start!r} in TAB 05')

def build(xlsx, songs_path):
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    tab02 = {r['Song ID'] for r in rows(tab(wb, 'TAB 02')) if r.get('Song ID')}
    site = {s['id'] for s in json.load(open(songs_path))['songs']}
    out, errors = [], []
    for r in rows(tab(wb, 'TAB 05')):
        eid = col(r, 'Expression ID')
        if not ID_RE.match(eid):
            continue
        songs = SONG_RE.findall(col(r, 'Linked song IDs'))
        for s in songs:
            if s not in tab02: errors.append(f'{eid}: {s} not in TAB 02')
            if s not in site: errors.append(f'{eid}: {s} not in data/songs.json')
        if songs and col(r, 'Match basis').strip() in ('', '—'):
            errors.append(f'{eid}: linked without a TAB 02 match basis')
        conf = col(r, 'Confidence').strip().lower()
        if conf not in CONF: errors.append(f'{eid}: unknown confidence {conf!r}')
        phrases = [col(r, 'Expression (as people say it)')] + [p.strip() for p in col(r, 'Also matches').split('|') if p.strip()]
        out.append({'id': eid, 'phrases': phrases, 'songs': songs, 'conf': conf,
                    'register': col(r, 'Register'), 'theme': col(r, 'Theme'),
                    'source': 'documented' if col(r, 'Source type').lower().startswith('documented') else 'ai-generated'})
    if errors:
        raise SystemExit('TAB 05 validation failed:\n  ' + '\n  '.join(errors))
    return {'schemaVersion': 1, 'generated': datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'),
            'source': 'SONIC REMEDY — MASTER SONG INDEX / TAB 05 — HUMAN EXPRESSION LIBRARY', 'count': len(out), 'expressions': out}

if __name__ == '__main__':
    if len(sys.argv) < 2: raise SystemExit(__doc__)
    dest = sys.argv[2] if len(sys.argv) > 2 else 'data/expressions.json'
    songs = sys.argv[3] if len(sys.argv) > 3 else 'data/songs.json'
    data = build(sys.argv[1], songs)
    json.dump(data, open(dest, 'w'), ensure_ascii=False, indent=1)
    print(f"{data['count']} expressions -> {dest} ({sum(1 for e in data['expressions'] if e['songs'])} linked)")
