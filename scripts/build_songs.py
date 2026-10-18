#!/usr/bin/env python3
"""Build data/songs.json for the Sonic Remedy hub discovery feature.

Source of truth: the master Sheet 'SONIC REMEDY — MASTER SONG INDEX'
(File > Download > Microsoft Excel, or a Drive export of the same file).

Usage:  python3 scripts/build_songs.py <master.xlsx> [data/songs.json]

Public data only. Excluded: Drive links, source-document references, working
notes (AI-inferred association notes, verification text), unpublished Episodes
17-20, and raw lyric / spoken-script text. Lyrics and artist scripts are reduced
to an unordered bag of content words used for matching only (never displayed).
Evidence distinctions are kept per field in `ev` (artist-confirmed |
artist-script | source-supported | ai-inferred) plus `conf` (sheet confidence).
"""
import sys, json, re, datetime
from collections import OrderedDict
from openpyxl import load_workbook

ID_RE = re.compile(r'^SR-F(\d{2})-T(\d{2})$')
# Formula-level states, verbatim from the hub's own formula tiles (index.html, artist-approved copy).
FORMULA_STATE = {1: 'CONFLICTED?', 2: 'SECOND-GUESSING?', 3: 'PUTTING IT OFF?', 4: 'INVISIBLE ROPES?',
                 5: 'HOLDING BACK?', 6: "DON'T FIT IN?", 7: "CAN'T SWITCH OFF?", 8: 'NEED TO RECONNECT?'}
FORMULA_URL = {7: 'https://sonicremedy07.vercel.app', 8: 'https://sonicremedy08.vercel.app'}
# Pairs/groups whose meaning overlaps strongly (artist-confirmed shared meanings and
# answer-song pairs).
OVERLAP_GROUPS = {
    'shay': ['SR-F05-T02', 'SR-F05-T03'],            # artist-confirmed shared meaning
    'city-of-gold': ['SR-F05-T04', 'SR-F05-T05'],    # artist-confirmed shared meaning
    'bet': ['SR-F02-T08', 'SR-F02-T09'],             # answer song to Track 08
    # 'progress' (SR-F07-T12 / T13) removed 2026-10-18: similar themes do not make compositions duplicates
}
# Site-config exclusions (artist-confirmed in TAB 02 col T): not intended for these visitors.
NOT_FOR = {
    'SR-F07-T07': ['betray', 'cheated on me', 'someone hurt me', 'they hurt me', 'he hurt me', 'she hurt me',
                   'broke my trust', 'lied to me', 'can\'t trust anyone', 'cant trust anyone', 'trust issues',
                   'was hurt by', 'been hurt by', 'got hurt', 'stabbed me in the back', 'backstab'],
}
STOP = set("""a an the and or but to of in on at for with my me i im i'm is am are be been it its this that so just
really very you your yours yourself myself we our us they them their he she his her not no nor do does did have has had
from as by was were will would can could should into out up down over than then there here what when where who how why
which while about all any some more most such only own same too s t don't can't won't it's you're they're i've i'll
let let's oh ooh yeah la na hey gonna wanna gotta got get go like one""".split())

def rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else '' for h in next(it)]
    for r in it:
        if r and any(c not in (None, '') for c in r):
            yield {h: ('' if v is None else str(v).strip()) for h, v in zip(head, r)}

def tab(wb, prefix):
    for n in wb.sheetnames:
        if n.upper().startswith(prefix):
            return wb[n]
    raise SystemExit(f'missing tab {prefix}')

NOTE_RE = re.compile(r'\s*\((?=[^)]*(?:card|doc\b|lyrics-doc|note|script not yet|topic from title|cover|Drive|vision))[^)]*\)', re.I)

def clean(s):
    """Drop parenthetical source/working notes such as '(pattern card 01: ...)'."""
    return re.sub(r'\s{2,}', ' ', NOTE_RE.sub('', s or '')).strip()

def split(s):
    return [p.strip().rstrip('.') for p in re.split(r';', clean(s)) if p.strip()]

def bag(texts):
    words = set()
    for t in texts:
        for w in re.findall(r"[a-z']+", t.lower().replace('\u2019', "'")):
            w = w.strip("'")
            if len(w) > 2 and w not in STOP:
                words.add(w)
    return sorted(words)

def conf(s):
    m = re.match(r'(High|Medium-High|Medium|Low-Medium|Low)', s or '')
    return m.group(1).lower() if m else None

def evidence(i):
    """Per-field evidence labels from TAB 02 col K (Support type)."""
    k = i.get('Support type', '')
    if k.startswith('Artist-provided spoken script'):
        return {'meaning': 'artist-script', 'situations': 'artist-script', 'desired': 'artist-script', 'expressions': 'ai-inferred'}
    if k.startswith('Artist-confirmed format, intent and list'):
        return {'meaning': 'artist-confirmed', 'situations': 'artist-confirmed', 'desired': 'artist-confirmed', 'expressions': 'ai-inferred'}
    if k.startswith('Artist-confirmed format'):   # F07 without script: topic from title + cover
        return {'meaning': 'ai-inferred', 'situations': 'ai-inferred', 'desired': 'ai-inferred', 'expressions': 'ai-inferred'}
    if 'aligne' in k or 'Shared meaning' in i.get('Verification status', ''):
        return {'meaning': 'artist-confirmed', 'situations': 'ai-inferred', 'desired': 'ai-inferred', 'expressions': 'ai-inferred'}
    if 'Source-supported' in k:
        return {'meaning': 'source-supported', 'situations': 'ai-inferred', 'desired': 'ai-inferred', 'expressions': 'ai-inferred'}
    return {'meaning': 'ai-inferred', 'situations': 'ai-inferred', 'desired': 'ai-inferred', 'expressions': 'ai-inferred'}

def comp_type(i, has_lyrics):
    m = (i.get('Composition type') or '').lower()
    if m.startswith('spoken'): return 'spoken-sleep'
    if 'wordless' in m: return 'instrumental-wordless'
    if m.startswith('instrumental'): return 'instrumental'
    return 'song'

def experience(ct, i):
    txt = ' '.join([i.get('Listening intentions', ''), i.get('Broader emotional states', ''), i.get('Composition type', ''), i.get('Intended listening experience', '')]).lower()
    tags = set()
    if ct == 'spoken-sleep': tags |= {'spoken', 'sleep', 'calm', 'night'}
    if ct.startswith('instrumental'): tags |= {'instrumental', 'wordless'}
    if ct == 'song': tags |= {'vocal'}
    if ct == 'song' and re.search(r'celebrat|danc|party|energetic|upbeat|playful|jazz-edm', txt): tags.add('energetic')
    if re.search(r'calm|peace|rest\b|stillness|soothing|quiet|tender', txt): tags.add('calm')
    return sorted(tags)

def build(xlsx, why_path):
    why = json.load(open(why_path))
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    master = list(rows(tab(wb, 'TAB 01')))
    intel = {r['Song ID']: r for r in rows(tab(wb, 'TAB 02'))}
    lyr, spoken, themes = {}, {}, {}
    for r in rows(tab(wb, 'TAB 03')):
        st = r.get('Verification status', '')
        if st.startswith('Verified'):
            lyr.setdefault(r['Song ID'], []).append(r['Exact lyric excerpt'])
        elif st.startswith('Artist-provided') and 'spoken-script' in st:
            spoken.setdefault(r['Song ID'], []).append(r['Exact lyric excerpt'])
        else:
            continue
        if r.get('Related theme or situation'):
            themes.setdefault(r['Song ID'], []).append(r['Related theme or situation'])
    group_of = {sid: g for g, ids in OVERLAP_GROUPS.items() for sid in ids}
    out = []
    for m in master:
        sid = m['Song ID']
        mm = ID_RE.match(sid)
        if not mm: continue
        fn, tn = int(mm.group(1)), int(mm.group(2))
        i = intel.get(sid, {})
        url = m['Individual song-page URL']
        if url.startswith('https://'): dest, dtype = url, 'song'
        else: dest, dtype = FORMULA_URL.get(fn), 'formula'
        ct = comp_type(i, bool(lyr.get(sid)))
        ev = evidence(i)
        if (i.get('Source evidence') or '').lower().startswith('artist-confirmed') and ev['situations'] == 'ai-inferred':
            ev['situations'] = ev['desired'] = 'artist-confirmed+ai-inferred'   # O/P artist-confirmed, D AI-inferred
        w = why[sid]
        # union of the song's own circumstances (D) and match-mode-A situations (O); for F07/F08
        # these are identical, for F06 O adds the artist-confirmed formula-wide association
        current = []
        for e in split(i.get('Specific circumstances where it may resonate', '')) + split(i.get('Current feeling or situation (match mode A)', '')):
            if e not in current: current.append(e)
        expressions = split(i.get('Alternative expressions people may use', ''))
        for e in split(i.get('Alternative everyday expressions', '')):
            if e not in expressions: expressions.append(e)
        out.append(OrderedDict([
            ('id', sid),
            ('title', m['Exact original song title']),
            ('formula', OrderedDict([('number', fn), ('name', m['Formula name']), ('state', FORMULA_STATE[fn])])),
            ('track', tn),
            ('url', dest),
            ('urlType', dtype),
            ('cover', f'assets/covers/{sid.lower()}.webp'),
            ('type', ct),
            ('experience', experience(ct, i)),
            ('why', w[0]),
            ('group', group_of.get(sid)),
            ('notFor', NOT_FOR.get(sid, [])),
            ('related', split(i.get('Related song IDs', ''))),
            ('m', OrderedDict([
                ('meaning', clean(i.get('Central meaning', ''))),
                ('intended', clean(i.get('Intended listening experience', ''))),
                ('patterns', split(i.get('Human patterns explored', ''))),
                ('situations', current),
                ('expressions', expressions),
                ('questions', [clean(i.get('Relevant questions or tensions', ''))] if i.get('Relevant questions or tensions') else []),
                ('desired', split(i.get('Desired state / felt need (match mode B)', ''))),
                ('moods', split(i.get('Broader emotional states', ''))),
                ('intentions', split(i.get('Listening intentions', ''))),
                ('themes', themes.get(sid, [])),
                ('lyricTerms', bag(lyr.get(sid, []))),
                ('spokenTerms', bag(spoken.get(sid, []))),
            ])),
            ('ev', OrderedDict([('why', w[1])] + list(ev.items()))),
            ('conf', conf(i.get('Confidence / verification status', ''))),
        ]))
    out.sort(key=lambda s: (s['formula']['number'], s['track']))
    return out

def validate(songs):
    errs = []
    ids = [s['id'] for s in songs]
    if len(songs) != 70: errs.append(f'expected 70 songs, got {len(songs)}')
    if len(ids) != len(set(ids)): errs.append('duplicate ids')
    for s in songs:
        if not s['title']: errs.append(f"{s['id']} no title")
        if not s['url'] or not re.match(r'^https://[a-z0-9-]+\.vercel\.app$', s['url']): errs.append(f"{s['id']} bad url {s['url']}")
        if not s['why']: errs.append(f"{s['id']} no why-line")
        if not s['m']['meaning']: errs.append(f"{s['id']} no meaning")
        for r in s['related']:
            if r not in ids: errs.append(f"{s['id']} unknown related {r}")
    blob = json.dumps(songs, ensure_ascii=False)
    for bad in ('drive.google.com', 'docs.google.com', 'EPISODE 17', 'Episode 17', 'Episodes 17', 'NOT VERIFIED'):
        if bad in blob: errs.append(f'forbidden text in public JSON: {bad}')
    return errs

if __name__ == '__main__':
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else 'data/songs.json'
    import os
    songs = build(src, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'why_lines.json'))
    errs = validate(songs)
    if errs:
        print('\n'.join(errs)); sys.exit(1)
    doc = OrderedDict([('schemaVersion', 3), ('generated', datetime.date.today().isoformat()),
                       ('source', 'SONIC REMEDY — MASTER SONG INDEX'), ('count', len(songs)), ('songs', songs)])
    with open(dst, 'w') as f:
        json.dump(doc, f, ensure_ascii=False, separators=(',', ':'))
    json.load(open(dst))
    print(f'OK {len(songs)} songs -> {dst} ({os.path.getsize(dst)//1024} KB)')
